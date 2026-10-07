# -*- coding: utf-8 -*-
"""
Document Extractor Service - Extração de texto de documentos.

Suporta:
- PDF (PyPDF2)
- DOCX (python-docx)
- TXT/MD (texto puro)

Pode baixar arquivos de URLs ou processar bytes diretamente.
"""

import io
import httpx
from typing import Optional, Dict, Any, Tuple
from pathlib import Path

# PDF extraction
try:
    import PyPDF2
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

# DOCX extraction
try:
    import docx
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False


class DocumentExtractor:
    """
    Extrai texto de documentos em diferentes formatos.

    Suporta download de URLs e processamento de bytes.
    """

    def __init__(self, timeout: float = 60.0):
        """
        Inicializa o extrator.

        Args:
            timeout: Timeout para downloads em segundos
        """
        self.timeout = timeout
        self._http_client: Optional[httpx.AsyncClient] = None

        # Log suporte
        print(f"[DOC-EXTRACTOR] PDF: {'OK' if PDF_SUPPORT else 'Não disponível'}")
        print(f"[DOC-EXTRACTOR] DOCX: {'OK' if DOCX_SUPPORT else 'Não disponível'}")

    async def _get_client(self) -> httpx.AsyncClient:
        """Retorna cliente HTTP (lazy init)."""
        if self._http_client is None:
            self._http_client = httpx.AsyncClient(timeout=self.timeout)
        return self._http_client

    async def download_file(self, url: str) -> Optional[bytes]:
        """
        Baixa arquivo de uma URL.

        Args:
            url: URL do arquivo

        Returns:
            Conteúdo em bytes ou None se falhar
        """
        try:
            client = await self._get_client()
            response = await client.get(url)
            response.raise_for_status()

            print(f"[DOC-EXTRACTOR] Download OK: {len(response.content)} bytes")
            return response.content

        except Exception as e:
            print(f"[DOC-EXTRACTOR] Erro no download: {e}")
            return None

    def extract_from_pdf(self, content: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extrai texto de PDF.

        Args:
            content: Bytes do PDF

        Returns:
            Tupla (texto, metadados)
        """
        if not PDF_SUPPORT:
            return "", {"error": "PyPDF2 não instalado"}

        try:
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            text_parts = []

            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text and text.strip():
                    text_parts.append(text.strip())

            full_text = "\n\n".join(text_parts)

            metadata = {
                "pages": len(reader.pages),
                "chars": len(full_text),
                "format": "PDF"
            }

            print(f"[DOC-EXTRACTOR] PDF: {metadata['pages']} páginas, {metadata['chars']} chars")
            return full_text, metadata

        except Exception as e:
            print(f"[DOC-EXTRACTOR] Erro PDF: {e}")
            return "", {"error": str(e)}

    def extract_from_docx(self, content: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extrai texto de DOCX.

        Args:
            content: Bytes do DOCX

        Returns:
            Tupla (texto, metadados)
        """
        if not DOCX_SUPPORT:
            return "", {"error": "python-docx não instalado"}

        try:
            doc = docx.Document(io.BytesIO(content))
            text_parts = []

            for para in doc.paragraphs:
                if para.text and para.text.strip():
                    text_parts.append(para.text.strip())

            # Também extrair de tabelas
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text and cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        text_parts.append(" | ".join(row_text))

            full_text = "\n\n".join(text_parts)

            metadata = {
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables),
                "chars": len(full_text),
                "format": "DOCX"
            }

            print(f"[DOC-EXTRACTOR] DOCX: {metadata['paragraphs']} parágrafos, {metadata['chars']} chars")
            return full_text, metadata

        except Exception as e:
            print(f"[DOC-EXTRACTOR] Erro DOCX: {e}")
            return "", {"error": str(e)}

    def extract_from_text(self, content: bytes) -> Tuple[str, Dict[str, Any]]:
        """
        Extrai texto de arquivo TXT/MD.

        Args:
            content: Bytes do arquivo

        Returns:
            Tupla (texto, metadados)
        """
        try:
            # Tentar UTF-8 primeiro
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                # Fallback para Latin-1
                text = content.decode("latin-1")

            metadata = {
                "chars": len(text),
                "lines": text.count("\n") + 1,
                "format": "TXT"
            }

            print(f"[DOC-EXTRACTOR] TXT: {metadata['lines']} linhas, {metadata['chars']} chars")
            return text, metadata

        except Exception as e:
            print(f"[DOC-EXTRACTOR] Erro TXT: {e}")
            return "", {"error": str(e)}

    def extract(self, content: bytes, file_type: str) -> Tuple[str, Dict[str, Any]]:
        """
        Extrai texto baseado no tipo do arquivo.

        Args:
            content: Bytes do arquivo
            file_type: Tipo (PDF, DOCX, TXT, etc)

        Returns:
            Tupla (texto, metadados)
        """
        file_type_upper = file_type.upper()

        if file_type_upper == "PDF":
            return self.extract_from_pdf(content)
        elif file_type_upper in ("DOCX", "DOC"):
            return self.extract_from_docx(content)
        elif file_type_upper in ("TXT", "TEXT", "MD", "MARKDOWN"):
            return self.extract_from_text(content)
        else:
            print(f"[DOC-EXTRACTOR] Tipo não suportado: {file_type}")
            return "", {"error": f"Tipo não suportado: {file_type}"}

    async def extract_from_url(self, url: str, file_type: str) -> Tuple[str, Dict[str, Any]]:
        """
        Baixa e extrai texto de uma URL.

        Args:
            url: URL do arquivo
            file_type: Tipo do arquivo

        Returns:
            Tupla (texto, metadados)
        """
        content = await self.download_file(url)
        if not content:
            return "", {"error": "Falha no download"}

        return self.extract(content, file_type)

    async def close(self):
        """Fecha cliente HTTP."""
        if self._http_client:
            await self._http_client.aclose()
            self._http_client = None


# Instância global
document_extractor: Optional[DocumentExtractor] = None
