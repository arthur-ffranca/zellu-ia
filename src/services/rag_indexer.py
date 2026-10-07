# -*- coding: utf-8 -*-
"""
RAG Indexer Service - Indexação de documentos no Pinecone.

Este serviço gerencia a indexação de documentos da base de conhecimento
das empresas no Pinecone para busca semântica (RAG).

Funcionalidades:
- Indexar documentos novos (create)
- Remover documentos deletados (delete)
- Atualizar documentos modificados (update)
- Buscar documentos similares (search)
- Gerenciar namespaces por empresa + knowledge_base

Atualizado em 10/12/2025:
- Namespace agora inclui knowledge_base_id para isolamento
- Integração com API Zellu para download e status updates
"""

import hashlib
import time
import uuid
from typing import List, Dict, Any, Optional
from datetime import datetime
from dataclasses import dataclass, field
from src.services.text_chunks import TextDocument as Document
from src.services.text_chunks import TextChunker
from src.services.embedding_cache import CachedOpenAIEmbeddings as OpenAIEmbeddings
from pinecone import Pinecone

from config import Settings
from src.services.rag_metrics_tracker import track_rag_metrics


@dataclass
class IndexedDocument:
    """Metadados de um documento indexado."""
    file_id: str
    company_id: str
    knowledge_base_id: Optional[str] = None  # NOVO: ID da base de conhecimento
    knowledge_base_name: Optional[str] = None  # NOVO: Nome da base
    file_name: str = ""
    file_type: str = ""
    title: Optional[str] = None
    description: Optional[str] = None
    folder: Optional[str] = None
    minio_path: Optional[str] = None
    content_hash: Optional[str] = None
    chunk_count: int = 0
    indexed_at: Optional[str] = None
    vector_ids: List[str] = field(default_factory=list)


class RAGIndexer:
    """
    Serviço de indexação RAG para documentos de empresas.

    Cada empresa tem seu próprio namespace no Pinecone para isolamento.
    Os documentos são divididos em chunks e indexados com metadados ricos.
    """

    def __init__(self, settings: Settings):
        """
        Inicializa o indexador RAG.

        Args:
            settings: Configurações da aplicação
        """
        self.settings = settings
        self.index_name = settings.PINECONE_INDEX_NAME

        # Inicializar embeddings
        # Nota: native runtime mais recente usa 'api_key' ao invés de 'openai_api_key'
        self.embeddings = OpenAIEmbeddings(
            api_key=settings.OPENAI_API_KEY,
            model=settings.OPENAI_EMBEDDING_MODEL
        )

        # Inicializar Pinecone
        self.pc = None
        self.index = None
        self._init_pinecone()

        # Text splitter para chunking
        self.text_splitter = TextChunker(
            chunk_size=settings.RAG_CHUNK_SIZE,
            chunk_overlap=settings.RAG_CHUNK_OVERLAP,
            separators=["\n\n", "\n", ". ", ", ", " ", ""],
            length_function=len
        )

        # Cache de documentos indexados por empresa
        # {company_id: {file_id: IndexedDocument}}
        self._indexed_docs: Dict[str, Dict[str, IndexedDocument]] = {}

        print(f"[RAG-INDEXER] Inicializado - Index: {self.index_name}")

    def _init_pinecone(self):
        """Inicializa conexão com Pinecone."""
        if not self.settings.PINECONE_API_KEY:
            print("[RAG-INDEXER] AVISO: PINECONE_API_KEY não configurado")
            return

        try:
            self.pc = Pinecone(api_key=self.settings.PINECONE_API_KEY)
            self.index = self.pc.Index(self.settings.PINECONE_INDEX_NAME)

            stats = self.index.describe_index_stats()
            print(f"[RAG-INDEXER] Conectado ao Pinecone - Vetores: {stats.get('total_vector_count', 0)}")

        except Exception as e:
            print(f"[RAG-INDEXER] Erro ao conectar Pinecone: {e}")
            self.index = None

    def _get_namespace(self, company_id: str, knowledge_base_id: Optional[str] = None) -> str:
        """
        Retorna o namespace do Pinecone para uma empresa/knowledge base.

        Args:
            company_id: ID da empresa
            knowledge_base_id: ID da base de conhecimento (opcional)

        Returns:
            Namespace no formato 'company_{id}' ou 'company_{id}_kb_{kb_id}'
        """
        if knowledge_base_id:
            return f"company_{company_id}_kb_{knowledge_base_id}"
        return f"company_{company_id}"

    def _compute_content_hash(self, content: str) -> str:
        """
        Calcula hash do conteúdo para detectar mudanças.

        Args:
            content: Conteúdo do documento

        Returns:
            Hash MD5 do conteúdo
        """
        return hashlib.md5(content.encode()).hexdigest()

    def _generate_vector_id(self, company_id: str, file_id: str, chunk_index: int) -> str:
        """
        Gera ID único para um vetor no Pinecone.

        Args:
            company_id: ID da empresa
            file_id: ID do arquivo
            chunk_index: Índice do chunk

        Returns:
            ID único no formato 'company_file_chunk'
        """
        return f"{company_id}_{file_id}_{chunk_index}"

    async def index_document(
        self,
        company_id: str,
        file_id: str,
        file_name: str,
        file_type: str,
        content: str,
        title: Optional[str] = None,
        description: Optional[str] = None,
        folder: Optional[str] = None,
        minio_path: Optional[str] = None,
        knowledge_base_id: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Indexa um documento no Pinecone.

        Args:
            company_id: ID da empresa
            file_id: ID do arquivo
            file_name: Nome do arquivo
            file_type: Tipo do arquivo (PDF, DOCX, TXT)
            content: Conteúdo textual do documento
            title: Título opcional
            description: Descrição opcional
            folder: Pasta/categoria opcional
            minio_path: Path no MinIO
            knowledge_base_id: ID da base de conhecimento (novo)

        Returns:
            Resultado da indexação com estatísticas
        """
        if not self.index:
            return {"success": False, "error": "Pinecone não disponível"}

        if not content or not content.strip():
            return {"success": False, "error": "Conteúdo vazio"}

        try:
            namespace = self._get_namespace(company_id, knowledge_base_id)
            content_hash = self._compute_content_hash(content)
            t0 = time.perf_counter()

            # Verificar se já está indexado com mesmo conteúdo
            if company_id in self._indexed_docs:
                existing = self._indexed_docs[company_id].get(file_id)
                if existing and existing.content_hash == content_hash:
                    print(f"[RAG-INDEXER] Documento já indexado (mesmo conteúdo): {file_name}")
                    return {
                        "success": True,
                        "action": "skipped",
                        "reason": "content_unchanged",
                        "chunk_count": existing.chunk_count
                    }
                elif existing:
                    # Conteúdo mudou - remover versão antiga
                    await self.delete_document(company_id, file_id)

            # Criar documento native runtime
            doc = Document(
                page_content=content,
                metadata={
                    "company_id": company_id,
                    "knowledge_base_id": knowledge_base_id or "",
                    "file_id": file_id,
                    "file_name": file_name,
                    "file_type": file_type,
                    "title": title or file_name,
                    "description": description or "",
                    "folder": folder or "",
                    "minio_path": minio_path or "",
                    "indexed_at": datetime.now().isoformat()
                }
            )

            # Dividir em chunks
            chunks = self.text_splitter.split_documents([doc])

            if not chunks:
                return {"success": False, "error": "Nenhum chunk gerado"}

            # Gerar embeddings e preparar vetores
            vectors = []
            vector_ids = []

            for i, chunk in enumerate(chunks):
                vector_id = self._generate_vector_id(company_id, file_id, i)
                vector_ids.append(vector_id)

                # Gerar embedding
                embedding = self.embeddings.embed_query(chunk.page_content)

                vectors.append({
                    "id": vector_id,
                    "values": embedding,
                    "metadata": {
                        "text": chunk.page_content,
                        "chunk_index": i,
                        "total_chunks": len(chunks),
                        **chunk.metadata
                    }
                })

            # Upsert em batches
            batch_size = 100
            for i in range(0, len(vectors), batch_size):
                batch = vectors[i:i + batch_size]
                self.index.upsert(vectors=batch, namespace=namespace)

            # Registrar documento indexado
            indexed_doc = IndexedDocument(
                file_id=file_id,
                company_id=company_id,
                file_name=file_name,
                file_type=file_type,
                title=title,
                description=description,
                folder=folder,
                minio_path=minio_path,
                content_hash=content_hash,
                chunk_count=len(chunks),
                indexed_at=datetime.now().isoformat(),
                vector_ids=vector_ids
            )

            if company_id not in self._indexed_docs:
                self._indexed_docs[company_id] = {}
            self._indexed_docs[company_id][file_id] = indexed_doc

            print(f"[RAG-INDEXER] Documento indexado: {file_name} ({len(chunks)} chunks)")

            track_rag_metrics({
                "operation": "indexing",
                "channel": "company_knowledge_base",
                "companyId": company_id,
                "knowledgeFileId": file_id,
                "eventId": str(uuid.uuid4()),
                "provider": "openai",
                "embeddingModel": self.settings.OPENAI_EMBEDDING_MODEL,
                "vectorStore": "pinecone",
                "indexNamespace": namespace,
                "chunkCount": len(chunks),
                "vectorsUpserted": len(vectors),
                "embeddingDimensions": len(vectors[0]["values"]) if vectors else None,
                "totalLatencyMs": int((time.perf_counter() - t0) * 1000),
                "status": "success",
                "metadata": {
                    "file_name": file_name,
                    "file_type": file_type,
                    "knowledge_base_id": knowledge_base_id or "",
                    "content_hash": content_hash,
                },
            })

            return {
                "success": True,
                "action": "indexed",
                "file_id": file_id,
                "file_name": file_name,
                "chunk_count": len(chunks),
                "namespace": namespace
            }

        except Exception as e:
            print(f"[RAG-INDEXER] Erro ao indexar {file_name}: {e}")
            return {"success": False, "error": str(e)}

    async def delete_document(
        self,
        company_id: str,
        file_id: str,
        knowledge_base_id: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Remove um documento do Pinecone.

        Args:
            company_id: ID da empresa
            file_id: ID do arquivo
            knowledge_base_id: ID da base de conhecimento (opcional)

        Returns:
            Resultado da remoção
        """
        if not self.index:
            return {"success": False, "error": "Pinecone não disponível"}

        try:
            namespace = self._get_namespace(company_id, knowledge_base_id)

            # Buscar IDs dos vetores do documento
            vector_ids = []
            if company_id in self._indexed_docs and file_id in self._indexed_docs[company_id]:
                vector_ids = self._indexed_docs[company_id][file_id].vector_ids

            if not vector_ids:
                # Tentar buscar por metadata (fallback)
                # Deletar por prefixo do ID
                prefix = f"{company_id}_{file_id}_"
                # Pinecone não suporta delete por prefixo diretamente
                # Precisamos listar e deletar

                # Por segurança, vamos deletar até 1000 chunks possíveis
                vector_ids = [f"{prefix}{i}" for i in range(1000)]

            # Deletar vetores
            if vector_ids:
                self.index.delete(ids=vector_ids, namespace=namespace)

            # Remover do cache
            if company_id in self._indexed_docs and file_id in self._indexed_docs[company_id]:
                del self._indexed_docs[company_id][file_id]

            print(f"[RAG-INDEXER] Documento removido: {file_id}")

            return {
                "success": True,
                "action": "deleted",
                "file_id": file_id,
                "vectors_removed": len(vector_ids)
            }

        except Exception as e:
            print(f"[RAG-INDEXER] Erro ao remover {file_id}: {e}")
            return {"success": False, "error": str(e)}

    async def search(
        self,
        company_id: str,
        query: str,
        top_k: int = 5,
        filter_folder: Optional[str] = None
    ) -> List[Dict[str, Any]]:
        """
        Busca documentos similares no Pinecone.

        Args:
            company_id: ID da empresa
            query: Texto da busca
            top_k: Número de resultados
            filter_folder: Filtrar por pasta (opcional)

        Returns:
            Lista de documentos relevantes com scores
        """
        if not self.index:
            return []

        try:
            namespace = self._get_namespace(company_id)

            # Gerar embedding da query
            query_embedding = self.embeddings.embed_query(query)

            # Construir filtro
            filter_dict = {"company_id": company_id}
            if filter_folder:
                filter_dict["folder"] = filter_folder

            # Buscar no Pinecone
            results = self.index.query(
                vector=query_embedding,
                top_k=top_k,
                include_metadata=True,
                namespace=namespace,
                filter=filter_dict
            )

            # Formatar resultados
            documents = []
            for match in results.get("matches", []):
                metadata = match.get("metadata", {})
                documents.append({
                    "content": metadata.get("text", ""),
                    "file_name": metadata.get("file_name", ""),
                    "title": metadata.get("title", ""),
                    "file_id": metadata.get("file_id", ""),
                    "folder": metadata.get("folder", ""),
                    "score": match.get("score", 0.0),
                    "chunk_index": metadata.get("chunk_index", 0),
                    "total_chunks": metadata.get("total_chunks", 1)
                })

            print(f"[RAG-INDEXER] Busca retornou {len(documents)} resultados")
            return documents

        except Exception as e:
            print(f"[RAG-INDEXER] Erro na busca: {e}")
            return []

    def get_company_stats(self, company_id: str) -> Dict[str, Any]:
        """
        Retorna estatísticas de indexação de uma empresa.

        Args:
            company_id: ID da empresa

        Returns:
            Estatísticas (docs, chunks, etc)
        """
        if company_id not in self._indexed_docs:
            return {
                "company_id": company_id,
                "total_documents": 0,
                "total_chunks": 0,
                "documents": []
            }

        docs = self._indexed_docs[company_id]
        total_chunks = sum(d.chunk_count for d in docs.values())

        return {
            "company_id": company_id,
            "total_documents": len(docs),
            "total_chunks": total_chunks,
            "documents": [
                {
                    "file_id": d.file_id,
                    "file_name": d.file_name,
                    "title": d.title,
                    "chunk_count": d.chunk_count,
                    "indexed_at": d.indexed_at
                }
                for d in docs.values()
            ]
        }

    def get_all_stats(self) -> Dict[str, Any]:
        """
        Retorna estatísticas globais de indexação.

        Returns:
            Estatísticas de todas as empresas
        """
        if not self.index:
            return {"error": "Pinecone não disponível"}

        try:
            stats = self.index.describe_index_stats()

            return {
                "total_vectors": stats.get("total_vector_count", 0),
                "dimension": stats.get("dimension", 0),
                "namespaces": stats.get("namespaces", {}),
                "companies_indexed": len(self._indexed_docs),
                "total_documents": sum(len(docs) for docs in self._indexed_docs.values())
            }
        except Exception as e:
            return {"error": str(e)}

    # =========================================================================
    # NOVO: Processamento de arquivos do webhook (10/12/2025)
    # =========================================================================

    async def process_webhook_files(
        self,
        company_id: str,
        knowledge_base_id: str,
        knowledge_base_name: str,
        files: List[Dict[str, Any]],
        api_client,  # ZelluAPIClient
        action: str = "create"
    ) -> Dict[str, Any]:
        """
        Processa arquivos recebidos via webhook de knowledge base.

        Para action='create':
        1. Baixa cada arquivo via API Zellu
        2. Extrai texto do arquivo
        3. Indexa no Pinecone
        4. Atualiza status na Zellu

        Para action='delete':
        1. Remove vetores do Pinecone

        Args:
            company_id: ID da empresa
            knowledge_base_id: ID da base de conhecimento
            knowledge_base_name: Nome da base
            files: Lista de arquivos do webhook
            api_client: Cliente da API Zellu
            action: 'create' ou 'delete'

        Returns:
            Resultado do processamento
        """
        results = {
            "company_id": company_id,
            "knowledge_base_id": knowledge_base_id,
            "action": action,
            "processed": [],
            "failed": [],
            "total_files": len(files)
        }

        namespace = self._get_namespace(company_id, knowledge_base_id)
        print(f"[RAG-INDEXER] Processando {len(files)} arquivos - namespace: {namespace}")

        for file_info in files:
            file_id = file_info.get("id")
            file_name = file_info.get("fileName", "")
            file_type = file_info.get("fileType", "")

            try:
                if action == "create":
                    # Atualizar status para in_progress (nao bloqueia se falhar)
                    try:
                        await api_client.update_file_processing_status(
                            file_id=file_id,
                            status="in_progress"
                        )
                    except Exception as status_err:
                        print(f"[RAG-INDEXER] Aviso: falha ao atualizar status para in_progress: {status_err}")

                    # Baixar arquivo
                    file_content = await api_client.download_knowledge_file(file_id)

                    if not file_content:
                        raise Exception("Falha ao baixar arquivo")

                    # Extrair texto do arquivo
                    text_content = await self._extract_text(file_content, file_type, file_name)

                    if not text_content:
                        resolved = self._resolve_file_type(file_type, file_name)
                        raise Exception(f"Nao foi possivel extrair texto do arquivo (type={file_type}, resolved={resolved}, name={file_name})")

                    # Indexar no Pinecone
                    result = await self.index_document(
                        company_id=company_id,
                        file_id=file_id,
                        file_name=file_name,
                        file_type=file_type,
                        content=text_content,
                        title=file_info.get("title"),
                        description=file_info.get("description"),
                        folder=file_info.get("folder"),
                        minio_path=file_info.get("minioPath"),
                        knowledge_base_id=knowledge_base_id
                    )

                    if result.get("success"):
                        # Atualizar status para completed
                        await api_client.update_file_processing_status(
                            file_id=file_id,
                            status="completed",
                            total_chunks=result.get("chunk_count", 0),
                            chunks_completed=result.get("chunk_count", 0),
                            chunks_failed=0
                        )
                        results["processed"].append({
                            "file_id": file_id,
                            "file_name": file_name,
                            "chunks": result.get("chunk_count", 0)
                        })
                    else:
                        raise Exception(result.get("error", "Erro desconhecido"))

                elif action == "delete":
                    # Remover do Pinecone
                    result = await self.delete_document(company_id, file_id, knowledge_base_id)
                    results["processed"].append({
                        "file_id": file_id,
                        "file_name": file_name,
                        "action": "deleted"
                    })

            except Exception as e:
                error_msg = str(e)
                print(f"[RAG-INDEXER] Erro processando {file_name}: {error_msg}")

                # Atualizar status para failed
                if action == "create":
                    await api_client.update_file_processing_status(
                        file_id=file_id,
                        status="failed",
                        error=error_msg
                    )

                results["failed"].append({
                    "file_id": file_id,
                    "file_name": file_name,
                    "error": error_msg
                })

        print(f"[RAG-INDEXER] Processamento concluido: {len(results['processed'])} ok, {len(results['failed'])} falhas")
        return results

    @staticmethod
    def _resolve_file_type(file_type: str, file_name: str) -> str:
        """Resolve o tipo de arquivo a partir do fileType ou da extensao do fileName."""
        # Mapa de MIME types para extensoes
        mime_map = {
            "application/pdf": "PDF",
            "application/msword": "DOC",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "DOCX",
            "application/vnd.ms-excel": "XLS",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "XLSX",
            "text/plain": "TXT",
            "text/csv": "CSV",
            "text/markdown": "MD",
            "application/json": "JSON",
        }

        # Tentar resolver MIME type
        ft_lower = file_type.strip().lower()
        if ft_lower in mime_map:
            resolved = mime_map[ft_lower]
            print(f"[RAG-INDEXER] fileType MIME '{file_type}' -> '{resolved}'")
            return resolved

        # Se ja e uma extensao curta (PDF, DOCX, TXT, etc.), usar direto
        ft_upper = file_type.strip().upper()
        if len(ft_upper) <= 5 and ft_upper.isalpha():
            return ft_upper

        # Fallback: extrair extensao do fileName
        if file_name and "." in file_name:
            ext = file_name.rsplit(".", 1)[-1].upper()
            print(f"[RAG-INDEXER] fileType '{file_type}' nao reconhecido, usando extensao do arquivo: '{ext}'")
            return ext

        return ft_upper

    async def _extract_text(self, content: bytes, file_type: str, file_name: str) -> Optional[str]:
        """
        Extrai texto de um arquivo binario.

        Args:
            content: Conteudo binario do arquivo
            file_type: Tipo do arquivo (extensao, MIME type, ou qualquer string)
            file_name: Nome do arquivo (usado como fallback para detectar tipo)

        Returns:
            Texto extraido ou None se falhar
        """
        resolved_type = self._resolve_file_type(file_type, file_name)
        print(f"[RAG-INDEXER] Extraindo texto: file_name={file_name}, file_type={file_type}, resolved={resolved_type}")

        try:
            if resolved_type in ("TXT", "MD", "JSON"):
                return content.decode("utf-8", errors="ignore")

            elif resolved_type == "PDF":
                import io
                try:
                    from PyPDF2 import PdfReader
                    reader = PdfReader(io.BytesIO(content))
                    text = ""
                    for page in reader.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
                    return text.strip()
                except ImportError:
                    print("[RAG-INDEXER] PyPDF2 nao instalado. Instale com: pip install PyPDF2")
                    return None

            elif resolved_type in ("DOCX", "DOC"):
                import io
                try:
                    from docx import Document as DocxDocument
                    doc = DocxDocument(io.BytesIO(content))
                    text = "\n".join([para.text for para in doc.paragraphs])
                    return text.strip()
                except ImportError:
                    print("[RAG-INDEXER] python-docx nao instalado. Instale com: pip install python-docx")
                    return None

            elif resolved_type in ("CSV", "XLS", "XLSX"):
                return content.decode("utf-8", errors="ignore")

            else:
                print(f"[RAG-INDEXER] Tipo de arquivo nao suportado: {file_type} (resolved: {resolved_type})")
                return None

        except Exception as e:
            print(f"[RAG-INDEXER] Erro ao extrair texto de {file_name}: {e}")
            return None

    async def index_document_with_kb(
        self,
        company_id: str,
        knowledge_base_id: str,
        file_id: str,
        file_name: str,
        file_type: str,
        content: str,
        **kwargs
    ) -> Dict[str, Any]:
        """
        Wrapper para index_document que inclui knowledge_base_id.
        """
        return await self.index_document(
            company_id=company_id,
            file_id=file_id,
            file_name=file_name,
            file_type=file_type,
            content=content,
            knowledge_base_id=knowledge_base_id,
            **kwargs
        )


# Instância global (inicializada em main.py)
rag_indexer: Optional[RAGIndexer] = None
